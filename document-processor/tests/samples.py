"""Builds small PDF, DOCX and Markdown files in memory, so the tests need no binary fixtures."""
import io

import docx
from fpdf import FPDF


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Aftercare Guide", level=0)  # "Title" style: skipped, the title comes from the file name
    document.add_heading("Tooth Extraction", level=1)
    document.add_paragraph("Bite gently on the gauze for 30 minutes.")
    document.add_heading("Eating and drinking", level=2)
    document.add_paragraph("Do not eat for two hours.", style="List Bullet")
    document.add_paragraph("Avoid hot drinks for 24 hours.", style="List Bullet")
    document.add_heading("Insurance", level=1)
    table = document.add_table(rows=2, cols=2)
    for row, (plan, covered) in enumerate([("Plan", "Covered"), ("Delta Dental PPO", "Yes")]):
        table.cell(row, 0).text, table.cell(row, 1).text = plan, covered
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_markdown() -> bytes:
    return b"""# Cancellation Policy

Please cancel at least 24 hours
in advance.

## Late cancellations

- A $50 fee applies.
- Two no-shows end online booking.
"""


def make_pdf() -> bytes:
    """Two pages: headings in larger fonts, a hyphenated word, page numbers,
    and a paragraph that runs from page 1 onto page 2."""
    pdf = FPDF(unit="pt", format="letter")
    pdf.set_auto_page_break(False)

    def line(text, size, gap=0):
        pdf.set_font("Helvetica", size=size)
        pdf.set_x(72)
        pdf.cell(0, size * 1.25, text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    pdf.add_page()
    pdf.set_y(72)
    line("Tooth Extraction", 20, gap=6)
    line("Bite gently on the gauze for thirty minutes after the", 11)
    line("procedure. Do not rinse your mouth during the first extrac-", 11)
    line("tion day.", 11, gap=10)
    line("Eating and drinking", 15, gap=4)
    line("Take the pain medication with food and", 11)
    pdf.set_y(740)
    line("Page 1", 9)

    pdf.add_page()
    pdf.set_y(72)
    line("plenty of water.", 11, gap=10)
    line("Insurance", 20, gap=6)
    line("We accept Delta Dental PPO plans.", 11)
    pdf.set_y(740)
    line("Page 2", 9)
    return bytes(pdf.output())


def make_blank_pdf() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    return bytes(pdf.output())
