"""Builds the demo documents (PDF and DOCX) from the Markdown sources in demo/documents/.

Run with the document processor's test image, which already has fpdf2 and python-docx:
    docker build --target test -t rag-document-tools document-processor
    docker run --rm -v "$PWD/demo:/demo" rag-document-tools python /demo/build_documents.py

Output: demo/build/<knowledge base>/<Title>.pdf | .docx | .md (the .md files are uploaded as Google Docs).
"""
import re
from pathlib import Path

import docx
from fpdf import FPDF

ROOT = Path(__file__).parent
SOURCE, OUTPUT = ROOT / "documents", ROOT / "build"
# Each format exercises a different extraction path
FORMATS = {
    "aftercare-guide": "pdf", "dental-emergencies": "pdf", "sterilization-protocol": "pdf",
    "insurance-and-payments": "docx", "time-off-and-scheduling": "docx",
    "cancellation-policy": "gdoc", "new-patient-faq": "gdoc", "staff-onboarding": "gdoc",
}
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
LIST_ITEM = re.compile(r"^(-|\d+\.)\s+(.*)$")


def parse(markdown: str):
    """Markdown subset used by the demo documents -> (kind, level or marker, text or rows)."""
    lines, items, i = markdown.splitlines(), [], 0
    while i < len(lines):
        line = lines[i]
        if heading := HEADING.match(line):
            items.append(("heading", len(heading[1]), heading[2]))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(set(c) <= {"-", ":"} for c in cells):
                    rows.append(cells)
                i += 1
            items.append(("table", None, rows))
            continue
        elif item := LIST_ITEM.match(line):
            items.append(("list", item[1], item[2]))
        elif line.strip():
            items.append(("paragraph", None, line.strip()))
        i += 1
    return items


class Pdf(FPDF):
    def footer(self):
        # Page numbers on purpose: the processor must drop them
        self.set_y(-40)
        self.set_font("Helvetica", size=9)
        self.cell(0, 12, f"Page {self.page_no()}", align="C")


def build_pdf(items, target: Path):
    pdf = Pdf(unit="pt", format="letter")
    pdf.set_margins(72, 72, 72)
    pdf.set_auto_page_break(True, margin=72)
    pdf.add_page()
    sizes = {1: 20, 2: 15, 3: 12.5}
    for kind, marker, text in items:
        if kind == "heading":
            size = sizes[marker]
            pdf.ln(size * 0.6)
            pdf.set_font("Helvetica", "B", size)
            pdf.multi_cell(0, size * 1.3, text, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(size * 0.4)
        else:
            pdf.set_font("Helvetica", size=11)
            prefix = "- " if marker == "-" else (f"{marker} " if marker else "")
            pdf.multi_cell(0, 11 * 1.3, prefix + text, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(11 * (0.3 if kind == "list" else 0.8))
    pdf.output(str(target))


def build_docx(items, target: Path):
    document = docx.Document()
    for kind, marker, content in items:
        if kind == "heading":
            document.add_heading(content, level=marker - 1)  # "#" is the document Title
        elif kind == "list":
            document.add_paragraph(content, style="List Bullet" if marker == "-" else "List Number")
        elif kind == "table":
            table = document.add_table(rows=len(content), cols=len(content[0]))
            table.style = "Table Grid"
            for r, row in enumerate(content):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
        else:
            document.add_paragraph(content)
    document.save(str(target))


def main():
    for source in sorted(SOURCE.glob("*/*.md")):
        markdown = source.read_text(encoding="utf-8")
        title = HEADING.match(markdown.splitlines()[0])[2]
        folder = OUTPUT / source.parent.name
        folder.mkdir(parents=True, exist_ok=True)
        kind = FORMATS[source.stem]
        if kind == "pdf":
            target = folder / f"{title}.pdf"
            build_pdf(parse(markdown), target)
        elif kind == "docx":
            target = folder / f"{title}.docx"
            build_docx(parse(markdown), target)
        else:
            target = folder / f"{title}.md"
            target.write_text(markdown, encoding="utf-8")
        print(f"{source.parent.name}/{target.name}")


if __name__ == "__main__":
    main()
